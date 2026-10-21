"""Reporting Service: звіти за період (FR-10, SR-10.1) та їх експорт у Excel, PDF і CSV.

Агрегати рахує БД (функції fn_report_uptime, fn_report_incidents, fn_report_ups_battery) —
звіт читає готові дані окремою транзакцією лише для читання, не блокуючи запис телеметрії.
"""
import csv
import io
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CATEGORY = {"ups": "ДБЖ", "switch": "Комутатор", "router": "Маршрутизатор", "onu": "ONU"}


@dataclass
class Report:
    key: str
    title: str
    headers: list[str]
    rows: list[list]


def _td(v) -> str:
    if v is None:
        return "—"
    s = int(v.total_seconds())
    return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"


def _n(v, digits=1):
    return None if v is None else round(float(v), digits)


async def build(db, p_from: datetime, p_to: datetime) -> list[Report]:
    up = await db.all("SELECT * FROM fn_report_uptime(%s, %s)", (p_from, p_to))
    inc = await db.all("SELECT * FROM fn_report_incidents(%s, %s)", (p_from, p_to))
    ups = await db.all("SELECT * FROM fn_report_ups_battery(%s, %s)", (p_from, p_to))
    return [
        Report("uptime", "Доступність обладнання",
               ["Пристрій", "Вузол", "Категорія", "Простій (год:хв:с)", "Доступність, %"],
               [[r["equipment"], r["location"], CATEGORY[r["category"]], _td(r["downtime"]), _n(r["uptime_pct"], 2)]
                for r in up]),
        Report("incidents", "Інциденти за вузлами та категоріями",
               ["Вузол", "Категорія", "Інцидентів", "Критичних", "Недоступність", "Відкриті", "Сер. MTTA", "Сер. MTTR"],
               [[r["location"], CATEGORY[r["category"]], r["incidents"], r["critical"], r["offline"], r["still_open"],
                 _td(r["avg_mtta"]), _td(r["avg_mttr"])] for r in inc]),
        Report("ups", "Стан акумуляторних батарей ДБЖ",
               ["ДБЖ", "Вузол", "Потужність, В·А", "Ємність АКБ, А·год", "Вік АКБ, діб", "Заряд зараз, %",
                "Мін. заряд, %", "Сер. t АКБ, °C", "Автономність, хв"],
               [[r["equipment"], r["location"], r["rated_power_va"], _n(r["battery_capacity_ah"]), r["battery_age_days"],
                 _n(r["charge_last"]), _n(r["charge_min"]), _n(r["temp_avg"]), _n(r["runtime_last_min"])] for r in ups]),
    ]


def period_title(p_from: datetime, p_to: datetime) -> str:
    return f"{p_from:%d.%m.%Y %H:%M} — {p_to:%d.%m.%Y %H:%M}"


def to_xlsx(reports: list[Report], period: str) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="1D3557")
    for rep in reports:
        ws = wb.create_sheet(rep.title[:31])
        ws.append([f"NetGuard Monitor — {rep.title}"])
        ws.append([f"Період: {period}"])
        ws.append([])
        ws.append(rep.headers)
        for c in ws[4]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), head_fill, Alignment(wrap_text=True)
        ws["A1"].font = Font(bold=True, size=13)
        for row in rep.rows:
            ws.append(row)
        for i, h in enumerate(rep.headers, 1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rep.rows]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 40)
        ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(rep: Report) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(rep.headers)
    w.writerows(rep.rows)
    return ("﻿" + buf.getvalue()).encode("utf-8")       # BOM — щоб Excel відкрив кирилицю


def to_pdf(reports: list[Report], period: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font_dir = Path("/usr/share/fonts/truetype/dejavu")
    if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVu", str(font_dir / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(font_dir / "DejaVuSans-Bold.ttf")))
    h1 = ParagraphStyle("h1", fontName="DejaVu-Bold", fontSize=14, spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName="DejaVu-Bold", fontSize=11, spaceBefore=10, spaceAfter=4)
    small = ParagraphStyle("s", fontName="DejaVu", fontSize=9, textColor=colors.grey)
    cell = ParagraphStyle("c", fontName="DejaVu", fontSize=8, leading=10)
    head = ParagraphStyle("hc", fontName="DejaVu-Bold", fontSize=8, leading=10, textColor=colors.white)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title="NetGuard Monitor — звіт")
    story = [Paragraph("NetGuard Monitor — звіт про роботу обладнання", h1),
             Paragraph(f"Період: {period}. Сформовано: {datetime.now():%d.%m.%Y %H:%M}", small)]
    for rep in reports:
        story.append(Paragraph(rep.title, h2))
        data = [[Paragraph(h, head) for h in rep.headers]] + \
               [[Paragraph("—" if v is None else str(v), cell) for v in r] for r in rep.rows]
        t = Table(data, repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1d3557")),
                               ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#b0b7c3")),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f4f7")]),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story += [t, Spacer(1, 4)]
    doc.build(story)
    return buf.getvalue()
